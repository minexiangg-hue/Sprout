#!/usr/bin/env python3
"""Export editable Markdown to print HTML, and rendered slides to a portable PPTX.
Requires: python -m pip install markdown python-pptx pypdf
First run: node scripts/export-deck.mjs (Playwright Chromium).
"""
from pathlib import Path
import re, json
import markdown
from pptx import Presentation
from pptx.util import Inches
ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/'docs'
body=markdown.markdown((DOCS/'business-plan.md').read_text(),extensions=['tables','fenced_code','toc'])
css='''@font-face{font-family:SproutSansSC;src:url('../apps/web/public/fonts/SproutSansSC.woff2')}*{box-sizing:border-box}body{font-family:SproutSansSC,system-ui,sans-serif;color:#26382d;font-size:11px;line-height:1.9;margin:0;background:#fbfcf8}main{max-width:1000px;margin:auto;padding:55px}h1{font-size:33px;line-height:1.45;border-bottom:3px solid #b5dc8c;padding-bottom:22px;margin:15px 0 26px}h2{font-size:20px;line-height:1.6;margin:34px 0 15px;break-after:avoid;color:#3b5c36}h3{break-after:avoid}p{margin:10px 0}a{color:#587949;overflow-wrap:anywhere}table{width:100%;border-collapse:collapse;font-size:10px;line-height:1.75;margin:18px 0}thead{display:table-header-group}tr{break-inside:avoid}th{background:#e5efda;text-align:left}th,td{border:1px solid #dbe4d1;padding:9px 10px;vertical-align:top}td:first-child{min-width:80px}code{font-family:SproutSansSC,system-ui,sans-serif;font-size:10px;word-break:break-word;background:#edf2e4;padding:2px 4px}pre{background:#edf2e4;padding:12px;white-space:pre-wrap}strong{font-weight:750}.coverline{font-size:10px;letter-spacing:3px;color:#849572}blockquote{border-left:3px solid #c0e98c;padding-left:18px}li{margin:5px 0}@page{size:A4;margin:18mm 16mm 20mm}@media print{body{background:white}main{padding:0}a{color:#365c3d}h2{break-after:avoid}p{orphans:3;widows:3}}'''
(DOCS/'business-plan.html').write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>Sprout 芽芽工坊 · 商业计划</title><style>'+css+'</style><main><div class="coverline">SPROUT · BUSINESS PLAN · 2026.09</div>'+body+'</main></html>')
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
notes=(DOCS/'pitch-content.md').read_text()
parts=re.split(r'^## (\d\d) / ',notes,flags=re.M)
page_notes={int(parts[i]):parts[i+1].split('\n## ')[0].strip() for i in range(1,len(parts),2)}
for n in range(1,15):
 slide=prs.slides.add_slide(prs.slide_layouts[6])
 slide.shapes.add_picture(str(DOCS/'slides'/f'{n:02}.png'),0,0,width=prs.slide_width,height=prs.slide_height)
 slide.notes_slide.notes_text_frame.text=page_notes.get(n,'')+'\n\n视觉保真演示版。设计与内容可编辑源：docs/pitch-deck.html、docs/pitch-content.md。'
prs.core_properties.title='Sprout · 芽芽工坊'
prs.core_properties.subject='把点子种出来，把每一步看明白。'
prs.core_properties.author='Sprout / GraphCode team'
prs.save(DOCS/'pitch-deck.pptx')
# A reusable native SVG asset, without raster conversion or third-party artwork.
html=(DOCS/'pitch-deck.html').read_text()
symbol=re.search(r'<symbol id="yaya" viewBox="0 0 240 240">(.*?)</symbol>',html,re.S).group(1)
assets=DOCS/'assets';assets.mkdir(exist_ok=True)
(assets/'yaya.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 240 240" role="img" aria-label="芽芽">'+symbol+'</svg>')
print('Created business-plan.html, pitch-deck.pptx (14 slides with speaker notes), assets/yaya.svg')
