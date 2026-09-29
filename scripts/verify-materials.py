#!/usr/bin/env python3
"""Verify exported handoff artifacts; requires pypdf and python-pptx."""
from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
import re,json,unicodedata
root=Path(__file__).resolve().parents[1]
normalize=lambda s:re.sub(r'\s+','',unicodedata.normalize('NFKC',s))
result={'documents':[]}
for relative,expected in [('docs/pitch-deck.pdf',14),('docs/business-plan.pdf',None)]:
 pdf=PdfReader(root/relative);text=normalize(''.join(page.extract_text() for page in pdf.pages))
 item={'file':relative,'pages':len(pdf.pages),'readableChinese': '芽芽工坊' in text,'noMissingGlyphs': '\x00' not in text,'pageSizesPt':sorted(set((round(float(p.mediabox.width)),round(float(p.mediabox.height))) for p in pdf.pages))}
 assert item['readableChinese'],relative
 assert item['noMissingGlyphs'],relative
 if expected:assert item['pages']==expected
 if relative.endswith('business-plan.pdf'):assert '单位经济' in text and '竞争' in text and '600,000' in text
 result['documents'].append(item)
prs=Presentation(root/'docs/pitch-deck.pptx')
assert len(prs.slides)==14
assert all(s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
result['powerpoint']={'slides':14,'speakerNotes':True,'format':'visual-fidelity image slides; editable source HTML/Markdown supplied'}
(root/'reports/materials-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(result,ensure_ascii=False,indent=2))
